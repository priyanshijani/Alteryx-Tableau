"""
generate_raw_data.py
--------------------
Generates a SYNTHETIC, deliberately messy CRM dataset for a fictional European
IT consultancy ("Nexora Consulting"). Output mimics what a Sales Ops analyst
really receives: monthly CRM extracts with a schema change after a CRM
migration, a stage-change audit log, an Excel account export with junk header
rows, an HR-style sales rep list, plus clean finance reference data.

All companies, people and numbers are fictional.

Run:  python scripts/generate_raw_data.py      (from the repo root)
Needs: pandas, numpy, openpyxl
"""
import os, random, datetime as dt
import numpy as np
import pandas as pd

SEED = 42
rng = np.random.default_rng(SEED)
random.seed(SEED)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "data", "raw", "crm_exports")
OPP_DIR = os.path.join(RAW, "opportunities")
REF = os.path.join(ROOT, "data", "reference")
for d in (RAW, OPP_DIR, REF):
    os.makedirs(d, exist_ok=True)

SNAPSHOT = dt.datetime(2026, 9, 15, 6, 0)        # CRM export time
START = dt.date(2024, 10, 1)                      # first monthly extract (CRM go-live)
SIM_START = dt.date(2024, 4, 1)                   # go-live also imported deals created Apr-Sep 2024
END = dt.date(2026, 9, 12)
MIGRATION_MONTH = (2025, 7)                       # CRM v2 schema from July 2025
N_DEALS = 4800          # ~4,000 created in the 24-month window + go-live import
DOMAIN = "nexora-consulting.eu"

# --------------------------------------------------------------------------
# Geography
# --------------------------------------------------------------------------
COUNTRIES = {
 # code: name, region, currency, weight, legal suffixes, raw variants, cities
 "DE": ("Germany", "DACH", "EUR", 18, ["GmbH", "AG", "SE"], ["Germany", "DE", "Deutschland", "germany", "GER"],
        [("Berlin", 52.520, 13.405), ("Munich", 48.137, 11.575), ("Hamburg", 53.551, 9.993),
         ("Frankfurt", 50.110, 8.682), ("Stuttgart", 48.776, 9.183), ("Düsseldorf", 51.227, 6.773)]),
 "AT": ("Austria", "DACH", "EUR", 3, ["GmbH", "AG"], ["Austria", "AT", "Österreich"],
        [("Vienna", 48.208, 16.373), ("Graz", 47.070, 15.439)]),
 "CH": ("Switzerland", "DACH", "CHF", 6, ["AG", "SA"], ["Switzerland", "CH", "Schweiz", "Suisse"],
        [("Zurich", 47.376, 8.541), ("Geneva", 46.204, 6.143), ("Basel", 47.559, 7.588)]),
 "NL": ("Netherlands", "Benelux", "EUR", 8, ["B.V.", "N.V."], ["Netherlands", "NL", "The Netherlands", "Holland"],
        [("Amsterdam", 52.370, 4.895), ("Rotterdam", 51.924, 4.477), ("Eindhoven", 51.441, 5.470)]),
 "BE": ("Belgium", "Benelux", "EUR", 3, ["NV", "SA"], ["Belgium", "BE", "Belgique", "België"],
        [("Brussels", 50.850, 4.352), ("Antwerp", 51.219, 4.402)]),
 "LU": ("Luxembourg", "Benelux", "EUR", 1, ["S.A.", "S.à r.l."], ["Luxembourg", "LU", "Lux"],
        [("Luxembourg", 49.611, 6.131)]),
 "GB": ("United Kingdom", "UK & Ireland", "GBP", 14, ["Ltd", "plc"], ["United Kingdom", "UK", "GB", "Great Britain", "England"],
        [("London", 51.507, -0.128), ("Manchester", 53.481, -2.243), ("Edinburgh", 55.953, -3.188), ("Birmingham", 52.486, -1.890)]),
 "IE": ("Ireland", "UK & Ireland", "EUR", 3, ["Ltd", "DAC"], ["Ireland", "IE", "Eire"],
        [("Dublin", 53.350, -6.260), ("Cork", 51.898, -8.475)]),
 "FR": ("France", "France", "EUR", 10, ["SA", "SAS"], ["France", "FR", "france"],
        [("Paris", 48.857, 2.352), ("Lyon", 45.764, 4.836), ("Toulouse", 43.605, 1.444), ("Lille", 50.629, 3.057)]),
 "SE": ("Sweden", "Nordics", "SEK", 5, ["AB"], ["Sweden", "SE", "Sverige"],
        [("Stockholm", 59.329, 18.069), ("Gothenburg", 57.709, 11.975)]),
 "DK": ("Denmark", "Nordics", "DKK", 3, ["A/S", "ApS"], ["Denmark", "DK", "Danmark"],
        [("Copenhagen", 55.676, 12.568), ("Aarhus", 56.163, 10.204)]),
 "NO": ("Norway", "Nordics", "NOK", 3, ["AS", "ASA"], ["Norway", "NO", "Norge"],
        [("Oslo", 59.914, 10.752), ("Bergen", 60.391, 5.322)]),
 "FI": ("Finland", "Nordics", "EUR", 2, ["Oy", "Oyj"], ["Finland", "FI", "Suomi"],
        [("Helsinki", 60.170, 24.938), ("Espoo", 60.205, 24.656)]),
 "ES": ("Spain", "Southern Europe", "EUR", 6, ["S.A.", "S.L."], ["Spain", "ES", "España", "Espana"],
        [("Madrid", 40.417, -3.704), ("Barcelona", 41.385, 2.173)]),
 "IT": ("Italy", "Southern Europe", "EUR", 7, ["S.p.A.", "S.r.l."], ["Italy", "IT", "Italia"],
        [("Milan", 45.464, 9.190), ("Rome", 41.903, 12.496), ("Turin", 45.070, 7.687)]),
 "PT": ("Portugal", "Southern Europe", "EUR", 2, ["S.A.", "Lda"], ["Portugal", "PT"],
        [("Lisbon", 38.722, -9.139), ("Porto", 41.158, -8.629)]),
 "PL": ("Poland", "CEE", "PLN", 5, ["S.A.", "Sp. z o.o."], ["Poland", "PL", "Polska"],
        [("Warsaw", 52.230, 21.012), ("Kraków", 50.065, 19.945), ("Wrocław", 51.108, 17.039)]),
 "CZ": ("Czechia", "CEE", "CZK", 2, ["a.s.", "s.r.o."], ["Czechia", "Czech Republic", "CZ"],
        [("Prague", 50.076, 14.438), ("Brno", 49.195, 16.607)]),
}
REGIONS = sorted({v[1] for v in COUNTRIES.values()})
REGION_CONV = {"Nordics": 1.05, "DACH": 1.02, "Benelux": 1.00, "UK & Ireland": 1.00,
               "France": 0.98, "Southern Europe": 0.95, "CEE": 1.00}
REGION_DUR = {"Nordics": 0.85, "DACH": 1.00, "Benelux": 0.95, "UK & Ireland": 0.95,
              "France": 1.10, "Southern Europe": 1.30, "CEE": 1.10}

# --------------------------------------------------------------------------
# Industries
# --------------------------------------------------------------------------
INDUSTRIES = {
 # name: raw variants, conversion mult, duration mult, name words
 "Financial Services": (["Financial Services", "FinServ", "Banking & Insurance", "financial services"], 1.00, 1.20,
                        ["Capital", "Bank", "Assurance", "Finance", "Invest"]),
 "Manufacturing & Automotive": (["Manufacturing & Automotive", "Manufacturing", "Automotive", "MFG"], 1.03, 1.00,
                        ["Motorenwerk", "Precision", "Industries", "Automotive", "Components"]),
 "Healthcare & Life Sciences": (["Healthcare & Life Sciences", "Healthcare", "Pharma", "HLS"], 1.00, 1.25,
                        ["Health", "Pharma", "Medical", "Biotech", "Clinics"]),
 "Retail & Consumer Goods": (["Retail & Consumer Goods", "Retail", "CPG", "retail"], 0.97, 0.95,
                        ["Retail", "Foods", "Brands", "Stores", "Consumer"]),
 "Energy & Utilities": (["Energy & Utilities", "Utilities", "Energy", "Energy/Utilities"], 1.06, 1.05,
                        ["Energy", "Power", "Grid", "Renewables", "Utilities"]),
 "Public Sector": (["Public Sector", "Government", "Public", "Gov"], 0.94, 1.50,
                        ["Agency", "Authority", "Municipal Services", "Public Services", "Institute"]),
 "Telecom & Media": (["Telecom & Media", "Telco", "Telecommunications", "Media"], 1.00, 1.00,
                        ["Telecom", "Media", "Networks", "Broadcast", "Connect"]),
 "Transport & Logistics": (["Transport & Logistics", "Logistics", "Transportation", "logistics"], 1.02, 1.00,
                        ["Logistics", "Freight", "Transport", "Shipping", "Mobility"]),
}
IND_NAMES = list(INDUSTRIES)
# Region-specific industry mix (energy transition in DACH/Nordics, FS in UK, etc.)
REGION_IND_WEIGHTS = {
 "DACH":            [14, 26, 12, 10, 16, 8, 7, 7],
 "Nordics":         [12, 14, 12, 8, 26, 10, 10, 8],
 "Benelux":         [16, 12, 16, 12, 12, 8, 10, 14],
 "UK & Ireland":    [28, 10, 14, 14, 8, 12, 9, 5],
 "France":          [16, 16, 14, 16, 10, 12, 9, 7],
 "Southern Europe": [16, 16, 10, 16, 10, 16, 10, 6],
 "CEE":             [18, 24, 8, 12, 12, 10, 10, 6],
}

SIZE_BANDS = ["<250", "250-999", "1,000-4,999", "5,000-19,999", "20,000+"]
SIZE_VARIANTS = {"<250": ["<250", "0-249", "small"], "250-999": ["250-999", "250 - 999"],
                 "1,000-4,999": ["1,000-4,999", "1000-4999", "1k-5k"],
                 "5,000-19,999": ["5,000-19,999", "5000-19999", "5k-20k"],
                 "20,000+": ["20,000+", "20000+", ">20k"]}
SIZE_MULT = {"<250": 0.4, "250-999": 0.7, "1,000-4,999": 1.0, "5,000-19,999": 1.5, "20,000+": 2.2}
SIZE_W = [12, 26, 32, 20, 10]

PREFIXES = ["Nordwind", "Alpenkraft", "Helvetia", "Brightwater", "Stellaris", "Rhein", "Baltica", "Lumen",
            "Aurora", "Castellan", "Meridian", "Vistula", "Danubia", "Iberia", "Solvik", "Fjordline",
            "Kestrel", "Oakridge", "Silverline", "Tamsin", "Valora", "Westbrook", "Zentra", "Arcadia",
            "Bergmann", "Corvina", "Delta", "Elbe", "Falkner", "Granite", "Harbour", "Isola",
            "Juniper", "Karlsson", "Lindholm", "Marisol", "Novara", "Orion", "Pionier", "Quadra",
            "Riviera", "Saxon", "Tauern", "Umbria", "Vektor", "Wexford", "Ysolde", "Zephyr",
            "Atlas", "Borealis", "Cygnus", "Dunmore", "Estrella", "Fontaine", "Gallway", "Hansa",
            "Invicta", "Jadeport", "Kronberg", "Lusitania", "Moselle", "Noord", "Olympia", "Polaris"]

# --------------------------------------------------------------------------
# Organisation
# --------------------------------------------------------------------------
DEPTS = {
 # name: code(v2), raw variants(v1), conversion mult, service lines {name: median EUR}
 "Cloud & Infrastructure": ("CLOUD", ["Cloud & Infrastructure", "Cloud & Infra", "Cloud and Infrastructure", "CLOUD"], 1.00,
     {"Cloud Migration": 180000, "Managed Cloud Services": 250000, "DevOps & Platform Engineering": 90000}),
 "Data & AI": ("DATA_AI", ["Data & AI", "Data and AI", "DATA & AI", "Data&AI"], 0.97,
     {"Data Platform Modernisation": 220000, "AI & GenAI Solutions": 150000, "Analytics & BI": 70000}),
 "Cybersecurity": ("CYBER", ["Cybersecurity", "Cyber Security", "Cyber", "cybersecurity"], 1.04,
     {"Security Assessment": 40000, "Managed SOC": 300000, "Identity & Access Management": 110000}),
 "ERP & Business Apps": ("ERP_APPS", ["ERP & Business Apps", "ERP & Apps", "ERP and Business Applications", "ERP"], 0.98,
     {"SAP S/4HANA Implementation": 600000, "Salesforce Implementation": 200000, "Application Support": 160000}),
}
DEPT_GROWTH = {"Cloud & Infrastructure": 0.3, "Data & AI": 1.6, "Cybersecurity": 0.5, "ERP & Business Apps": 0.0}
NC = ["DACH", "Benelux", "Nordics", "CEE"]           # North & Central cluster
WS = ["UK & Ireland", "France", "Southern Europe"]    # West & South cluster

MANAGERS = [
 # manager, dept, cluster, conversion skill, volume, optimism
 ("Claudia Weber", "Cloud & Infrastructure", NC, 1.03, 1.0, 0.0),
 ("James Whitaker", "Cloud & Infrastructure", WS, 0.98, 1.0, 0.0),
 ("Stefan Brandt", "Data & AI", NC, 1.02, 1.0, 0.05),
 ("Richard Evans", "Data & AI", WS, 0.88, 1.35, 0.45),   # big pipeline, low conversion, optimistic
 ("Henrik Dahl", "Cybersecurity", NC, 1.07, 0.9, 0.0),   # small team, great win rate
 ("Marie Lefèvre", "Cybersecurity", WS, 1.00, 1.0, 0.05),
 ("Katrin Vogel", "ERP & Business Apps", NC, 1.00, 0.9, 0.0),
 ("Giulia Romano", "ERP & Business Apps", WS, 0.97, 0.9, 0.1),
]
AES = [
 # name, manager, regions, hire, leave
 ("Lukas Schneider", "Claudia Weber", ["DACH"], "2021-03-01", None),
 ("Freja Nielsen", "Claudia Weber", ["Nordics"], "2022-01-10", None),
 ("Pieter Janssens", "Claudia Weber", ["Benelux"], "2020-09-01", None),
 ("Tomáš Novák", "Claudia Weber", ["CEE"], "2023-02-01", None),
 ("Oliver Hughes", "James Whitaker", ["UK & Ireland"], "2019-06-03", None),
 ("Chloé Martin", "James Whitaker", ["France"], "2021-11-15", None),
 ("Mateo López", "James Whitaker", ["Southern Europe"], "2022-04-04", None),
 ("Anna Müller", "Stefan Brandt", ["DACH"], "2020-02-01", None),
 ("Mikael Lindqvist", "Stefan Brandt", ["Nordics"], "2021-08-16", None),
 ("Emma de Vries", "Stefan Brandt", ["Benelux"], "2022-10-03", None),
 ("Katarzyna Nowak", "Stefan Brandt", ["CEE"], "2023-05-02", None),
 ("Ethan Clarke", "Richard Evans", ["UK & Ireland"], "2021-01-11", None),
 ("Hugo Laurent", "Richard Evans", ["France"], "2022-03-01", None),
 ("Marco Rossi", "Richard Evans", ["Southern Europe"], "2020-05-04", "2025-08-31"),
 ("Elena Bianchi", "Richard Evans", ["Southern Europe"], "2025-09-15", None),
 ("Niamh Walsh", "Richard Evans", ["UK & Ireland"], "2023-09-04", None),
 ("Felix Wagner", "Henrik Dahl", ["DACH"], "2020-07-01", None),
 ("Sara Andersson", "Henrik Dahl", ["Nordics", "Benelux"], "2021-04-06", None),
 ("Jan Kowalski", "Henrik Dahl", ["CEE"], "2022-06-01", None),
 ("Tom Barker", "Marie Lefèvre", ["UK & Ireland"], "2019-10-01", "2026-03-31"),
 ("Aoife Byrne", "Marie Lefèvre", ["UK & Ireland"], "2026-04-13", None),
 ("David Moreau", "Marie Lefèvre", ["France"], "2021-02-15", None),
 ("Isabel Ferreira", "Marie Lefèvre", ["Southern Europe"], "2022-09-12", None),
 ("Julia Hoffmann", "Katrin Vogel", ["DACH"], "2019-03-18", None),
 ("Sven Johansson", "Katrin Vogel", ["Nordics"], "2021-06-01", None),
 ("Lena Fischer", "Katrin Vogel", ["DACH", "Benelux", "CEE"], "2023-01-09", None),
 ("Paolo Conti", "Giulia Romano", ["Southern Europe"], "2020-11-02", None),
 ("Grace Thompson", "Giulia Romano", ["UK & Ireland"], "2022-02-14", None),
 ("Camille Bernard", "Giulia Romano", ["France"], "2021-09-06", None),
]
MGR = {m[0]: m for m in MANAGERS}


def ascii_fold(s):
    table = str.maketrans("äöüÄÖÜéèêëáàâíìîóòôúùûçñšžčřýěåøæłśźżńćęąŁŚŻ",
                          "aouAOUeeeeaaaiiiooouuucnszcryeaoalszzncea" + "LSZ")
    return s.translate(table).replace("ß", "ss")


def email_of(name):
    first, last = name.split(" ", 1)
    return f"{ascii_fold(first).lower()}.{ascii_fold(last).lower().replace(' ', '')}@{DOMAIN}"


def d(s):
    return dt.datetime.strptime(s, "%Y-%m-%d").date() if s else None


# --------------------------------------------------------------------------
# 1) Accounts
# --------------------------------------------------------------------------
codes = list(COUNTRIES)
cw = np.array([COUNTRIES[c][3] for c in codes], float); cw /= cw.sum()
accounts, used_names = [], set()
N_ACC = 720
for i in range(N_ACC):
    c = rng.choice(codes, p=cw)
    name_c, region, cur, _, suffixes, _, cities = COUNTRIES[c]
    iw = np.array(REGION_IND_WEIGHTS[region], float); iw /= iw.sum()
    ind = rng.choice(IND_NAMES, p=iw)
    while True:
        base = f"{random.choice(PREFIXES)} {random.choice(INDUSTRIES[ind][3])}"
        if base not in used_names:
            used_names.add(base); break
    size = rng.choice(SIZE_BANDS, p=np.array(SIZE_W) / sum(SIZE_W))
    city, lat, lon = random.choice(cities)
    rev_base = {"<250": 40, "250-999": 180, "1,000-4,999": 900, "5,000-19,999": 3500, "20,000+": 15000}[size]
    accounts.append(dict(account_id=f"ACC-{10001 + i}", name=f"{base} {random.choice(suffixes)}", base=base,
                         country=c, region=region, currency=cur, industry=ind, size=size, city=city,
                         lat=round(lat + rng.normal(0, 0.03), 5), lon=round(lon + rng.normal(0, 0.03), 5),
                         revenue=round(rev_base * rng.lognormal(0, 0.4), 1)))
acc_df = pd.DataFrame(accounts)

# duplicate account records (same company, typed differently)
dup_rows, dup_map = [], {}
dup_src = acc_df.sample(45, random_state=SEED)
for k, (_, a) in enumerate(dup_src.iterrows()):
    variant = random.choice([a["name"].upper(), a["name"].lower(), f"  {a['name']}", a["base"],
                             a["name"].replace(" ", "  ", 1), a["name"] + "."])
    new_id = f"ACC-{10001 + N_ACC + k}"
    dup_map[a["account_id"]] = new_id
    r = a.to_dict(); r.update(account_id=new_id, name=variant)
    dup_rows.append(r)
acc_all = pd.concat([acc_df, pd.DataFrame(dup_rows)], ignore_index=True)

# --------------------------------------------------------------------------
# 2) Reps
# --------------------------------------------------------------------------
reps = []
for name, mgr, regions, hire, leave in AES:
    _, dept, _, skill, vol, opt = MGR[mgr]
    reps.append(dict(name=name, email=email_of(name), manager=mgr, dept=dept, regions=regions,
                     hire=d(hire), leave=d(leave), skill=skill * rng.normal(1, 0.04), vol=vol, opt=opt))

# --------------------------------------------------------------------------
# 3) Simulate deals
# --------------------------------------------------------------------------
STAGES = {1: "Discovery", 2: "Solution Design", 3: "Proposal", 4: "Negotiation", 5: "Closed Won", 0: "Closed Lost"}
BASE_P = {1: 0.70, 2: 0.78, 3: 0.70, 4: 0.82}
BASE_DUR = {1: 28, 2: 32, 3: 30, 4: 22}
PROB_CHOICES = {1: [10, 20, 30, 40], 2: [50, 60, 70], 3: [75, 80, 85], 4: [90, 95], 5: [100], 0: [0]}
LOSS_EARLY = ["No Budget", "No Decision / Stalled", "Not a Fit"]
LOSS_LATE = ["Lost to Competitor", "Price", "Timing / Postponed", "Scope Reduced / In-house", "No Decision / Stalled"]
LEAD_SOURCES = ["Inbound - Website", "Partner Referral", "Outbound - SDR", "Event / Webinar", "Hyperscaler Co-sell"]


def active_window(r):
    a = max(SIM_START, r["hire"]); b = min(END, r["leave"] or END)
    return a, b


def sample_created(r):
    a, b = active_window(r)
    g = DEPT_GROWTH[r["dept"]]
    span = (b - a).days
    while True:
        x = rng.random()
        day = a + dt.timedelta(days=int(x * span))
        t = (day - SIM_START).days / (END - SIM_START).days
        if rng.random() < (1 + g * t) / (1 + g):
            return dt.datetime.combine(day, dt.time(int(rng.integers(8, 18)), int(rng.integers(0, 60)), int(rng.integers(0, 60))))


w = np.array([r["vol"] * max((active_window(r)[1] - active_window(r)[0]).days, 0) for r in reps], float)
alloc = np.floor(w / w.sum() * N_DEALS).astype(int)
alloc[np.argmax(alloc)] += N_DEALS - alloc.sum()

deals = []
for r, n in zip(reps, alloc):
    for _ in range(n):
        created = sample_created(r)
        region = random.choice(r["regions"])
        pool = acc_df[acc_df.region == region]
        sw = pool["size"].map(SIZE_MULT).values ** 0.7
        # Data & AI and Cloud lean on Energy & Manufacturing
        if r["dept"] in ("Data & AI", "Cloud & Infrastructure"):
            sw = sw * np.where(pool.industry.isin(["Energy & Utilities", "Manufacturing & Automotive"]), 1.6, 1.0)
        acc = pool.sample(1, weights=sw, random_state=int(rng.integers(1e9))).iloc[0]
        dept = r["dept"]
        sl_map = DEPTS[dept][3]
        sl = random.choice(list(sl_map))
        amount = sl_map[sl] * SIZE_MULT[acc["size"]] * rng.lognormal(0, 0.55)
        amount_eur = max(5000, round(amount / 500) * 500)
        ind_conv, ind_dur = INDUSTRIES[acc.industry][1], INDUSTRIES[acc.industry][2]
        conv = r["skill"] * DEPTS[dept][2] * REGION_CONV[region] * ind_conv
        durm = REGION_DUR[region] * ind_dur
        t, stage, hist, regress = created, 1, [(None, 1, created)], 0
        while True:
            mean = BASE_DUR[stage] * durm
            if dept == "Data & AI" and stage == 3:
                mean *= 1.8                      # GenAI proposals wait on budget approval
            if acc.industry == "Public Sector" and stage == 4:
                mean *= 1.8                      # procurement
            dur = rng.gamma(2.0, mean / 2.0)
            if rng.random() < 0.07:
                dur *= rng.uniform(2.5, 5.0)     # stuck deal
            t_next = t + dt.timedelta(days=float(dur), hours=float(rng.uniform(0, 8)))
            if t_next > SNAPSHOT - dt.timedelta(hours=12):
                break
            u = rng.random()
            if stage in (3, 4) and regress < 2 and u < 0.06:
                new = stage - 1; regress += 1
            elif rng.random() < min(0.95, BASE_P[stage] * conv):
                new = stage + 1
            else:
                new = 0
            hist.append((stage, new, t_next)); t, stage = t_next, new
            if stage in (0, 5):
                break
        closed = stage in (0, 5)
        close_dt = t if closed else None
        lost_at = hist[-1][0] if stage == 0 else None
        loss_reason = (random.choice(LOSS_EARLY) if lost_at in (1, 2) else random.choice(LOSS_LATE)) if stage == 0 else None
        # rep-entered probability
        if stage in (1, 2, 3, 4):
            ch = PROB_CHOICES[stage]
            prob = ch[-1] if rng.random() < r["opt"] else random.choice(ch)
            if rng.random() < 0.03 + r["opt"] * 0.35:
                prob = min(95, prob + random.choice([15, 20, 25]))   # out of band (optimistic)
        else:
            prob = PROB_CHOICES[stage][0]
        # expected close
        exp_close = (created + dt.timedelta(days=float(max(30, rng.normal(115, 35) * durm)))).date()
        if not closed and exp_close < SNAPSHOT.date() and rng.random() < 0.55:
            exp_close = SNAPSHOT.date() + dt.timedelta(days=int(rng.integers(10, 100)))
        # last activity
        stage_entered = hist[-1][2]
        if closed:
            last_act = close_dt.date()
        else:
            span = max((SNAPSHOT - stage_entered).days, 0)
            if rng.random() < 0.18:
                last_act = max(created.date(), (stage_entered - dt.timedelta(days=int(rng.integers(0, 10)))).date())
            else:
                last_act = (stage_entered + dt.timedelta(days=int(rng.integers(0, span + 1)))).date()
        last_mod = max(hist[-1][2], dt.datetime.combine(last_act, dt.time(17, 0)))
        last_mod = min(last_mod, SNAPSHOT - dt.timedelta(hours=1))
        deals.append(dict(created=created, owner=r, account=acc, region=region, dept=dept, service=sl,
                          amount_eur=amount_eur, stage=stage, hist=hist, close_dt=close_dt, loss=loss_reason,
                          prob=prob, exp_close=exp_close, last_act=last_act, last_mod=last_mod))

deals.sort(key=lambda x: x["created"])

# IDs with a few gaps (deleted opportunities)
all_ids = list(range(1, N_DEALS + 16))
deleted = sorted(random.sample(all_ids[50:-50], 15))
kept = [i for i in all_ids if i not in deleted][:N_DEALS]
for dl, i in zip(deals, kept):
    dl["num"] = i

# lead source: expansion if the account already won before
first_won = {}
for dl in sorted(deals, key=lambda x: x["close_dt"] or dt.datetime.max):
    if dl["stage"] == 5:
        first_won.setdefault(dl["account"].account_id, dl["close_dt"])
for dl in deals:
    fw = first_won.get(dl["account"].account_id)
    if fw and fw < dl["created"] and rng.random() < 0.6:
        dl["lead"] = "Existing Customer (Upsell)"
    else:
        dl["lead"] = random.choice(LEAD_SOURCES)

# ---- "stored" CRM values (data-entry problems live in the system itself) ----
leaver_open_keep = 0.85
for dl in deals:
    r = dl["owner"]; a = dl["account"]
    dl["acc_id_stored"] = a.account_id
    if a.account_id in dup_map and rng.random() < 0.5:
        dl["acc_id_stored"] = dup_map[a.account_id]
    dl["currency"] = a.currency if rng.random() > 0.08 else "EUR"
    dl["amount_stored"] = dl["amount_eur"]           # set below in local currency
    # open deals of leavers: some were never reassigned
    dl["owner_stored"] = r
    if r["leave"] and dl["stage"] in (1, 2, 3, 4):
        if rng.random() > leaver_open_keep:
            repl = [x for x in reps if x["manager"] == r["manager"] and x["leave"] is None and set(x["regions"]) & set(r["regions"])]
            if repl:
                dl["owner_stored"] = repl[0]

# FX: monthly rates (EUR per 1 unit)
months = pd.period_range("2024-10", "2026-09", freq="M")
fx_months = pd.period_range("2024-04", "2026-09", freq="M")
fx_base = {"EUR": 1.0, "GBP": 1.17, "CHF": 1.05, "SEK": 0.087, "DKK": 0.134, "NOK": 0.086, "PLN": 0.232, "CZK": 0.040}
fx_rows = []
for cur, b in fx_base.items():
    v = b
    for m in fx_months:
        if cur != "EUR":
            v = v * (1 + rng.normal(0, 0.012))
        fx_rows.append((str(m), cur, round(v, 5)))
fx = pd.DataFrame(fx_rows, columns=["Month", "Currency", "EUR per 1 unit"])
fx_lookup = {(m, c): v for m, c, v in fx_rows}

for dl in deals:
    m = dl["created"].strftime("%Y-%m")
    local = dl["amount_eur"] / fx_lookup[(m, dl["currency"])]
    dl["amount_stored"] = float(round(local / 100) * 100)

# stored-value problems
idx = list(range(len(deals)))
for i in random.sample([j for j in idx if deals[j]["amount_eur"] >= 60000], 3):
    deals[i]["amount_stored"] *= 100          # typo outlier
for i in random.sample(idx, 4):
    deals[i]["amount_stored"] *= -1           # negative
for i in idx:
    if deals[i]["stage"] == 1 and rng.random() < 0.12:
        deals[i]["amount_stored"] = None      # amount not yet estimated
    elif rng.random() < 0.004:
        deals[i]["amount_stored"] = None
for i in idx:
    if rng.random() < 0.03:
        deals[i]["currency"] = None           # currency not captured
for i in idx:                                  # won deals without close date entered
    if deals[i]["stage"] == 5 and rng.random() < 0.03:
        deals[i]["close_missing"] = True
orphans = random.sample(idx, 10)
for n, i in enumerate(orphans):
    deals[i]["acc_id_stored"] = f"ACC-19{n:03d}"   # account deleted from CRM

# --------------------------------------------------------------------------
# 4) Write monthly opportunity extracts (schema changes after migration)
# --------------------------------------------------------------------------
STAGE_V1 = {1: ["Discovery", "1 - Discovery", "Qualification", "discovery"],
            2: ["Solution Design", "2 - Solution Design", "Solutioning", "solution design"],
            3: ["Proposal", "3 - Proposal", "Proposal Sent", "proposal/quote"],
            4: ["Negotiation", "4 - Negotiation", "Contracting", "Negotiation/Review"],
            5: ["Closed Won", "Won", "5 - Closed Won", "closed won"],
            0: ["Closed Lost", "Lost", "Closed - Lost", "closed lost"]}
STAGE_V2 = {1: ["Stage 1: Discovery (0-50%)", "Discovery"], 2: ["Stage 2: Solution Design (50-75%)", "Solution Design"],
            3: ["Stage 3: Proposal (75-90%)", "Proposal"], 4: ["Stage 4: Negotiation (90-100%)", "Negotiation"],
            5: ["Stage 5: Closed Won (100%)", "Closed Won"], 0: ["Stage 0: Closed Lost", "Closed Lost"]}
LOSS_VAR = {"Lost to Competitor": ["Lost to Competitor", "Competitor", "lost to competitor"],
            "Price": ["Price", "PRICE", "Too expensive"], "No Budget": ["No Budget", "Budget", "no budget"],
            "No Decision / Stalled": ["No Decision / Stalled", "No decision", "Stalled"],
            "Not a Fit": ["Not a Fit", "Not qualified", "No fit"],
            "Timing / Postponed": ["Timing / Postponed", "Postponed", "Timing"],
            "Scope Reduced / In-house": ["Scope Reduced / In-house", "Doing in-house", "In-house"]}
LEAD_VAR = {"Inbound - Website": ["Inbound - Website", "Website", "Inbound"],
            "Partner Referral": ["Partner Referral", "Partner", "Referral"],
            "Outbound - SDR": ["Outbound - SDR", "Outbound", "SDR"],
            "Event / Webinar": ["Event / Webinar", "Event", "Webinar"],
            "Hyperscaler Co-sell": ["Hyperscaler Co-sell", "AWS/Azure Co-sell", "Co-sell"],
            "Existing Customer (Upsell)": ["Existing Customer (Upsell)", "Upsell", "Existing Customer"]}
SYMBOL = {"EUR": "€", "GBP": "£", "CHF": "CHF ", "SEK": "SEK ", "DKK": "DKK ", "NOK": "NOK ", "PLN": "PLN ", "CZK": "CZK "}
FCAT = {1: "Pipeline", 2: "Pipeline", 3: "Best Case", 4: "Commit", 5: "Closed", 0: "Omitted"}


def pick(lst, p_first=0.7):
    return lst[0] if rng.random() < p_first else random.choice(lst[1:]) if len(lst) > 1 else lst[0]


def fmt_date(x, v2):
    if x is None:
        return ""
    if isinstance(x, dt.datetime):
        x = x.date()
    u = rng.random()
    if not v2:
        return x.strftime("%d/%m/%Y") if u < 0.03 else x.isoformat()
    return x.isoformat() if u < 0.05 else x.strftime("%d.%m.%Y") if u < 0.07 else x.strftime("%d/%m/%Y")


def fmt_amount(a, cur):
    if a is None:
        return ""
    u = rng.random()
    if u < 0.72:
        return f"{a:.0f}"
    if u < 0.82:
        return f"{a:,.2f}"
    if u < 0.87:
        return f"{a:.2f}"
    if u < 0.94 and cur:
        return f"{SYMBOL[cur]}{a:,.0f}"
    return f"{a:,.0f}".replace(",", " ")


def fmt_currency(c):
    if c is None:
        return ""
    u = rng.random()
    return c.lower() if u < 0.05 else ("€" if c == "EUR" and u < 0.09 else c)


def owner_name_variant(name):
    first, last = name.split(" ", 1)
    u = rng.random()
    if u < 0.82: return name
    if u < 0.88: return name.lower()
    if u < 0.93: return f"{last}, {first}"
    if u < 0.97: return name + "  "
    return ascii_fold(name).replace("u", "ue", 1) if "ü" in name else f"{first[0]}. {last}"


def row_for(dl, v2, stale=False):
    stage = dl["stage"]; prob = dl["prob"]; lm = dl["last_mod"]; close = dl["close_dt"]; loss = dl["loss"]
    if stale:  # an older snapshot of the same opportunity
        prev = dl["hist"][-2]
        stage = prev[1]; prob = random.choice(PROB_CHOICES[stage]); close = None; loss = None
        lm = dl["hist"][-1][2] - dt.timedelta(days=float(rng.uniform(1, 5)))
    close_val = None if (dl.get("close_missing") or stage not in (0, 5)) else close
    owner = dl["owner_stored"]; acc = dl["account"]
    if not v2:
        acc_name = acc["name"] if rng.random() > 0.1 else acc["name"].upper()
        return {
            "Opportunity ID": f"OPP-{dl['num']}",
            "Opportunity Name": f"{acc['base']} - {dl['service']}",
            "Account ID": dl["acc_id_stored"],
            "Account Name": acc_name,
            "Opportunity Owner": owner_name_variant(owner["name"]),
            "Department": pick(DEPTS[dl["dept"]][1]),
            "Service Line": dl["service"] if rng.random() > 0.03 else "",
            "Stage": pick(STAGE_V1[stage], 0.55),
            "Probability (%)": (f"{prob}%" if rng.random() < 0.3 else str(prob)) if rng.random() > 0.02 else "",
            "Amount": fmt_amount(dl["amount_stored"], dl["currency"]),
            "Currency": fmt_currency(dl["currency"]),
            "Created Date": fmt_date(dl["created"], False),
            "Expected Close Date": fmt_date(dl["exp_close"], False),
            "Actual Close Date": fmt_date(close_val, False),
            "Loss Reason": (pick(LOSS_VAR[loss]) if rng.random() > 0.1 else "") if loss else "",
            "Lead Source": pick(LEAD_VAR[dl["lead"]]),
            "Last Activity Date": fmt_date(dl["last_act"], False),
            "Last Modified": lm.strftime("%Y-%m-%d %H:%M:%S"),
        }
    em = owner["email"]
    return {
        "Opp_ID": f"OPP-{dl['num']:06d}",
        "Created_On": fmt_date(dl["created"], True),
        "Opp_Name": f"{acc['base']} - {dl['service']}",
        "Acct_ID": dl["acc_id_stored"],
        "Owner_Email": em.upper() if rng.random() < 0.05 else em,
        "Business_Unit": DEPTS[dl["dept"]][0],
        "Offering": dl["service"] if rng.random() > 0.04 else "",
        "Sales_Stage": pick(STAGE_V2[stage], 0.8),
        "Win_Probability": (str(prob / 100) if rng.random() < 0.4 else str(prob)) if rng.random() > 0.02 else "",
        "Deal_Value": fmt_amount(dl["amount_stored"], dl["currency"]),
        "Deal_Currency": fmt_currency(dl["currency"]),
        "Close_Date_Expected": fmt_date(dl["exp_close"], True),
        "Close_Date_Actual": fmt_date(close_val, True),
        "Lost_Reason": (pick(LOSS_VAR[loss]) if rng.random() > 0.1 else "") if loss else "",
        "Source": pick(LEAD_VAR[dl["lead"]]),
        "Last_Activity": fmt_date(dl["last_act"], True),
        "Forecast_Category": FCAT[stage] if rng.random() > 0.15 else random.choice(["Commit", "Best Case", "Pipeline"]),
        "Modified_At": lm.strftime("%d/%m/%Y %H:%M"),
    }


test_rows_made = 0
file_stats = []
for m in months:
    y, mo = m.year, m.month
    v2 = (y, mo) >= MIGRATION_MONTH
    if (y, mo) == (START.year, START.month):   # first extract also carries the go-live import
        mdeals = [x for x in deals if x["created"].date() < dt.date(y, mo, 1) or (x["created"].year == y and x["created"].month == mo)]
    else:
        mdeals = [x for x in deals if x["created"].year == y and x["created"].month == mo]
    rows = []
    for dl in mdeals:
        rows.append(row_for(dl, v2))
        if len(dl["hist"]) >= 2 and rng.random() < 0.03:
            rows.append(row_for(dl, v2, stale=True))
        if rng.random() < 0.015:
            rows.append(dict(rows[-1]))
    if mo in (11, 3, 6, 9) and test_rows_made < 6:
        for k in range(2 if mo == 9 else 1):
            test_rows_made += 1
            fake = dict(rows[0])
            idkey = "Opp_ID" if v2 else "Opportunity ID"
            fake[idkey] = f"OPP-{990000 + test_rows_made:06d}" if v2 else f"OPP-{990000 + test_rows_made}"
            fake["Opp_Name" if v2 else "Opportunity Name"] = "TEST - please ignore"
            fake["Acct_ID" if v2 else "Account ID"] = "ACC-00000"
            if v2:
                fake["Owner_Email"] = f"crm.admin@{DOMAIN}"
            else:
                fake["Opportunity Owner"] = "CRM Admin"; fake["Account Name"] = "ACME Test Account"
            fake["Deal_Value" if v2 else "Amount"] = "1"
            rows.insert(int(rng.integers(0, len(rows))), fake)
    rng.shuffle(rows)
    df = pd.DataFrame(rows)
    fn = os.path.join(OPP_DIR, f"opps_{y}_{mo:02d}.csv")
    df.to_csv(fn, index=False, encoding="utf-8")
    file_stats.append((os.path.basename(fn), len(df)))

# --------------------------------------------------------------------------
# 5) Stage history audit log
# --------------------------------------------------------------------------
STAGE_LOG = {1: ["Discovery", "1 - Discovery", "Stage 1: Discovery (0-50%)"],
             2: ["Solution Design", "2 - Solution Design", "Stage 2: Solution Design (50-75%)"],
             3: ["Proposal", "3 - Proposal", "Stage 3: Proposal (75-90%)"],
             4: ["Negotiation", "4 - Negotiation", "Stage 4: Negotiation (90-100%)"],
             5: ["Closed Won", "Won", "Stage 5: Closed Won (100%)"],
             0: ["Closed Lost", "Lost", "Stage 0: Closed Lost"]}


def log_id(num):
    u = rng.random()
    if u < 0.90: return f"OPP-{num:06d}"
    if u < 0.95: return f"opp-{num:06d}"
    if u < 0.98: return f" OPP-{num:06d} "
    return f"OPP-{num}"


def log_ts(t):
    u = rng.random()
    if u < 0.6: return t.strftime("%Y-%m-%d %H:%M:%S")
    if u < 0.85: return t.strftime("%d/%m/%Y %H:%M")
    return t.strftime("%Y-%m-%dT%H:%M:%SZ")


log = []
for dl in deals:
    by = dl["owner"]["email"]
    for old, new, t in dl["hist"]:
        log.append([log_id(dl["num"]), "Stage", "" if old is None else pick(STAGE_LOG[old], 0.6),
                    pick(STAGE_LOG[new], 0.6), t, by if old is not None else "integration.user@" + DOMAIN])
    if rng.random() < 0.25:
        for _ in range(int(rng.integers(1, 3))):
            t = dl["created"] + dt.timedelta(days=float(rng.uniform(1, 40)))
            if t < SNAPSHOT:
                fld = random.choice(["Amount", "Close Date", "Next Step"])
                log.append([log_id(dl["num"]), fld, "", "(changed)", t, by])
for num in deleted:   # history of opportunities later deleted in CRM
    t0 = dt.datetime.combine(SIM_START, dt.time(10)) + dt.timedelta(days=int(rng.integers(0, 600)))
    log.append([f"OPP-{num:06d}", "Stage", "", "Discovery", t0, "integration.user@" + DOMAIN])
    log.append([f"OPP-{num:06d}", "Stage", "Discovery", "Solution Design", t0 + dt.timedelta(days=20), "system"])
for k in range(1, 7):
    t0 = dt.datetime(2025, 3, 3, 9) + dt.timedelta(days=30 * k)
    log.append([f"OPP-{990000 + k:06d}", "Stage", "", "Discovery", t0, f"crm.admin@{DOMAIN}"])
log.sort(key=lambda x: x[4])
dups = random.sample(range(len(log)), int(len(log) * 0.01))
for i in sorted(dups, reverse=True):
    log.insert(i, list(log[i]))
log_df = pd.DataFrame(log, columns=["Opportunity ID", "Field", "Old Value", "New Value", "Changed On", "Changed By"])
log_df["Changed On"] = log_df["Changed On"].apply(log_ts)
log_df.insert(0, "Log ID", [f"L{100000 + i}" for i in range(len(log_df))])
log_df.to_csv(os.path.join(RAW, "stage_history_log.csv"), index=False, encoding="utf-8")

# --------------------------------------------------------------------------
# 6) Accounts Excel export (junk header rows, variants, missing / swapped coords)
# --------------------------------------------------------------------------
from openpyxl import Workbook
from openpyxl.styles import Font
wb = Workbook(); ws = wb.active; ws.title = "Accounts"
ws["A1"] = "Report: All Accounts - Europe"; ws["A1"].font = Font(name="Arial", bold=True, size=12)
ws["A2"] = f"Exported by sales.ops@{DOMAIN} on 15/09/2026 06:00"
hdr = ["Account ID", "Account Name", "Industry", "Country", "City", "Employees", "Annual Revenue (EUR m)", "Latitude", "Longitude", "Website"]
ws.append([]); ws.append(hdr)
for c in ws[4]:
    c.font = Font(name="Arial", bold=True)
acc_out = acc_all.sample(frac=1, random_state=SEED)
for _, a in acc_out.iterrows():
    lat, lon = a["lat"], a["lon"]
    u = rng.random()
    if u < 0.04: lat, lon = None, None
    elif u < 0.047: lat, lon = lon, lat
    rev = a["revenue"] if rng.random() > 0.06 else random.choice(["n/a", None])
    city = a["city"] if rng.random() > 0.08 else random.choice([a["city"].upper(), a["city"].lower(), " " + a["city"]])
    web = "www." + ascii_fold(a["base"]).lower().replace(" ", "") + "." + (a["country"].lower() if a["country"] != "GB" else "co.uk")
    ws.append([a["account_id"], a["name"], pick(INDUSTRIES[a["industry"]][0], 0.72),
               pick(COUNTRIES[a["country"]][5], 0.72), city, pick(SIZE_VARIANTS[a["size"]], 0.75),
               rev, lat, lon, web])
ws.append([]); ws.append([f"Total records: {len(acc_out)}"]); ws.append(["Confidential - Internal Use Only"])
for col, wdt in zip("ABCDEFGHIJ", [12, 34, 26, 16, 14, 13, 22, 10, 10, 30]):
    ws.column_dimensions[col].width = wdt
wb.save(os.path.join(RAW, "accounts_export.xlsx"))

# --------------------------------------------------------------------------
# 7) Sales team roster (HR / Sales Ops file)
# --------------------------------------------------------------------------
roster = []
for i, (mname, dept, cluster, *_ ) in enumerate(MANAGERS):
    roster.append([f"E{2001 + i}", mname, email_of(mname), "Sales Manager", pick(DEPTS[dept][1], 0.6), "",
                   "North & Central" if cluster == NC else "West & South", "01/01/2019", "", ""])
for i, r in enumerate(reps):
    quota = {"Cloud & Infrastructure": 1_400_000, "Data & AI": 1_200_000, "Cybersecurity": 900_000,
             "ERP & Business Apps": 1_800_000}[r["dept"]]
    nm = r["name"] + (" " if rng.random() < 0.15 else "")
    roster.append([f"E{3001 + i}", nm, r["email"], "Account Executive", pick(DEPTS[r["dept"]][1], 0.6), r["manager"],
                   "; ".join(r["regions"]), r["hire"].strftime("%d/%m/%Y"),
                   r["leave"].strftime("%d/%m/%Y") if r["leave"] else "", quota])
pd.DataFrame(roster, columns=["Employee ID", "Full Name", "Email", "Role", "Department", "Sales Manager",
                              "Territory", "Hire Date", "Leave Date", "Annual Quota (EUR)"]) \
  .to_csv(os.path.join(RAW, "sales_team_roster.csv"), index=False, encoding="utf-8")

# --------------------------------------------------------------------------
# 8) Reference data (clean, owned by Finance / Sales Ops)
# --------------------------------------------------------------------------
fx.to_csv(os.path.join(REF, "fx_rates_monthly.csv"), index=False)

won = pd.DataFrame([(x["dept"], x["close_dt"], x["amount_eur"]) for x in deals if x["stage"] == 5],
                   columns=["dept", "close", "eur"])
won["q"] = pd.PeriodIndex(won["close"], freq="Q")
tgt_rows = []
quarters = pd.period_range("2024Q4", "2026Q4", freq="Q")
for dept in DEPTS:
    s_ = won[won.dept == dept].groupby("q")["eur"].sum().reindex(quarters[:-1]).fillna(0)
    x = np.arange(len(s_)); b1, b0 = np.polyfit(x[1:], s_.values[1:], 1)   # plan = trend + stretch
    for i, q in enumerate(quarters):
        val = max(b0 + b1 * i, s_.values[1:].mean()) * rng.uniform(1.03, 1.15)
        tgt_rows.append([dept, f"{q.year}-Q{q.quarter}", int(round(val / 10000) * 10000)])
pd.DataFrame(tgt_rows, columns=["Department", "Quarter", "Bookings Target (EUR)"]) \
  .to_csv(os.path.join(REF, "department_targets_quarterly.csv"), index=False)

mapping = [
 ("v1", "Opportunity ID", "opportunity_id"), ("v2", "Opp_ID", "opportunity_id"),
 ("v1", "Opportunity Name", "opportunity_name"), ("v2", "Opp_Name", "opportunity_name"),
 ("v1", "Account ID", "account_id"), ("v2", "Acct_ID", "account_id"),
 ("v1", "Account Name", "account_name"),
 ("v1", "Opportunity Owner", "owner_name"), ("v2", "Owner_Email", "owner_email"),
 ("v1", "Department", "department_raw"), ("v2", "Business_Unit", "department_raw"),
 ("v1", "Service Line", "service_line"), ("v2", "Offering", "service_line"),
 ("v1", "Stage", "stage_raw"), ("v2", "Sales_Stage", "stage_raw"),
 ("v1", "Probability (%)", "probability_raw"), ("v2", "Win_Probability", "probability_raw"),
 ("v1", "Amount", "amount_raw"), ("v2", "Deal_Value", "amount_raw"),
 ("v1", "Currency", "currency_raw"), ("v2", "Deal_Currency", "currency_raw"),
 ("v1", "Created Date", "created_date_raw"), ("v2", "Created_On", "created_date_raw"),
 ("v1", "Expected Close Date", "expected_close_raw"), ("v2", "Close_Date_Expected", "expected_close_raw"),
 ("v1", "Actual Close Date", "actual_close_raw"), ("v2", "Close_Date_Actual", "actual_close_raw"),
 ("v1", "Loss Reason", "loss_reason_raw"), ("v2", "Lost_Reason", "loss_reason_raw"),
 ("v1", "Lead Source", "lead_source_raw"), ("v2", "Source", "lead_source_raw"),
 ("v1", "Last Activity Date", "last_activity_raw"), ("v2", "Last_Activity", "last_activity_raw"),
 ("v2", "Forecast_Category", "forecast_category_rep"),
 ("v1", "Last Modified", "last_modified_raw"), ("v2", "Modified_At", "last_modified_raw"),
]
pd.DataFrame(mapping, columns=["crm_version", "source_field", "standard_field"]) \
  .to_csv(os.path.join(REF, "crm_field_mapping_v1_v2.csv"), index=False)

# CRM control totals (what the source system reports; used for reconciliation)
ct = []
all_ids_in_crm = [x["num"] for x in deals] + [990000 + k for k in range(1, test_rows_made + 1)]
ct.append(("Unique opportunities in CRM (incl. test records)", "All", len(all_ids_in_crm)))
for s, nm in STAGES.items():
    ct.append(("Opportunities by stage (excl. test records)", nm, sum(1 for x in deals if x["stage"] == s)))
cur_tot = {}
for x in deals:
    if x["amount_stored"] is not None:
        k = x["currency"] or "(blank)"
        cur_tot[k] = cur_tot.get(k, 0) + x["amount_stored"]
for k in sorted(cur_tot):
    ct.append(("Sum of Amount as stored, local currency (excl. test records)", k, round(cur_tot[k], 2)))
ct.append(("Opportunities with blank Amount (excl. test records)", "All", sum(1 for x in deals if x["amount_stored"] is None)))
ct.append(("Account records in CRM", "All", len(acc_all)))
ct_df = pd.DataFrame(ct, columns=["Control Total", "Dimension", "Value"])
ct_df["Value"] = ct_df["Value"].apply(lambda v: f"{int(v)}" if float(v).is_integer() else f"{v:.2f}")
ct_df.to_csv(os.path.join(REF, "crm_control_totals.csv"), index=False)

print("Opportunity extracts:")
for f, n in file_stats:
    print(f"  {f}: {n} rows")
print("Stage log rows:", len(log_df), "| Accounts:", len(acc_all), "| Test rows:", test_rows_made)
